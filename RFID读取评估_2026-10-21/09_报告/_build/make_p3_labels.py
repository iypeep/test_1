# -*- coding: utf-8 -*-
"""P4 表头改成「只读到 1 个 / 1–2 个 / 1–3 个 …」；P3 补小电子书 SN；大/小平板 → 大/小电子书

1) P4：≤1..≤5 的表头语义就是「读到 1 个」「读到 1–2 个」…（分母＝读到 ≥1 个的箱），
   把表头直接写成中文口径，避免被读成「含 0 个」的累计；脚注同步改写。
2) P3：小电子书 4 行补 SN（同样取自 EpcLog 真实读取；第 2 行读到 2 个，所以给两个 SN）。
3) P3：`大平板 / 小平板` → `大电子书 / 小电子书`（含现象/结论文字与副标题）。

真实 SN 来源（EpcLog_20261021.log，SETTINGNAME AX7907AAA036 同批）：
    大电子书 4 行：GN46HP01641502GU / GN46HP01642601PC / GN46HP01641603MC / GN46HP01642601JM
    小电子书 1 行：GN46HP01642601KX
             2 行：GN46HP01642601JR + GN46HP01641604EH   （该行读到 2 个）
             3 行：GN456V03643201X6
             4 行：GN456U02642707JL

用法：py -3.12 make_p3_labels.py --pptx "<deck>.pptx"
"""
import os, argparse, shutil, datetime

from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

LATIN, EA = "Arial", "微软雅黑"

HEAD_NEW = {6: "只读到\n1 个", 7: "1–2 个", 8: "1–3 个", 9: "1–4 个", 10: "1–5 个"}

SMALL_SNS = {
    5: "GN46HP01642601KX",
    6: "GN46HP01642601JR\nGN46HP01641604EH",
    7: "GN456V03643201X6",
    8: "GN456U02642707JL",
}

RENAME = [("大平板", "大电子书"), ("小平板", "小电子书"), ("大/小平板", "大/小电子书")]
# 副标题里的「原报告称大/小平板」是历史称谓，改回去
RESTORE = [("（原报告称大/小电子书）", "（原报告称大/小平板）")]
EXTRA = [("平板是金属件", "电子书背板是金属件")]

FOOT_OLD = ("分母＝该组「读到 ≥1 个」的箱（读到 0 个的箱不计入，只在「箱数 / 读到≥1 箱」两列体现）；"
            "≤N＝这些箱里读到 ≤N 个的占比（累计），读全 6 个＝读到 6 个。"
            "31 天 113,667 箱中可映射机型 83,706 箱，其中只有 5,247 箱（6.3%）读到过标签。")
FOOT_NEW = ("分母＝「读到 ≥1 个」的箱（读到 0 个的箱不计入，见「箱数 / 读到≥1 箱」两列）；"
            "「只读到 1 个」＝该箱只读到 1 个，「1–2 个」＝读到 1 个或 2 个，依次累计；读全 6 个＝读到 6 个。"
            "31 天 113,667 箱中可映射机型 83,706 箱，其中 5,247 箱（6.3%）读到过标签。")


def set_font(run, size, bold=None, color="1A1A1A"):
    f = run.font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def cell_text(cell, text, size=8.4, bold=False, color="1A1A1A", align=None):
    tf = cell.text_frame
    tf.paragraphs[0].clear()
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        if align is not None:
            p.alignment = align
        r = p.add_run(); r.text = line
        set_font(r, size, bold, color)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    a = ap.parse_args()
    src = a.pptx
    shutil.copy2(src, os.path.splitext(src)[0] + ".bak-p3lab-" +
                 datetime.datetime.now().strftime("%H%M%S") + ".pptx")
    prs = Presentation(src)

    # ---------------- P4：表头 + 脚注
    p4 = prs.slides[3]
    t = None
    for sh in p4.shapes:
        if getattr(sh, "has_table", False) and sh.has_table and sh.table.cell(0, 0).text.strip() == "HQ代码":
            t = sh.table
    if t is None:
        raise SystemExit("P4 表未找到")
    for ci, txt in HEAD_NEW.items():
        cell_text(t.cell(0, ci), txt, size=7.6, bold=True)
    print("P4 headers updated")
    for sh in p4.shapes:
        if sh.has_text_frame and ("分母＝" in sh.text_frame.text or "≤N＝" in sh.text_frame.text):
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.text = FOOT_NEW
            print("P4 footnote updated")

    # ---------------- P3：改名 + 补小电子书 SN
    p3 = prs.slides[2]
    cnt = 0
    for sh in p3.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    for old, new in RENAME:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                            cnt += 1
                    for old, new in RESTORE + EXTRA:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                            cnt += 1
    tb = None
    for sh in p3.shapes:
        if getattr(sh, "has_table", False) and sh.has_table and sh.table.cell(0, 8).text.strip().endswith("SN"):
            tb = sh.table
    if tb is None:
        raise SystemExit("P3 表未找到")
    for ri, txt in SMALL_SNS.items():
        cell_text(tb.cell(ri, 8), txt, size=8.0, align=PP_ALIGN.LEFT)
        cell_text(tb.cell(ri, 0), "小电子书", size=8.4)
    for ri in range(1, 5):
        cell_text(tb.cell(ri, 0), "大电子书", size=8.4)
    print(f"P3 renamed runs={cnt}; SN added for rows {sorted(SMALL_SNS)}")

    prs.save(src)
    print("saved", src, os.path.getsize(src))

    # 复核
    prs2 = Presentation(src)
    left = []
    for i, sl in enumerate(prs2.slides, 1):
        for sh in sl.shapes:
            txts = []
            if sh.has_text_frame:
                txts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                txts += [c.text for row in sh.table.rows for c in row.cells]
            for x in txts:
                if "大平板" in x or "小平板" in x:
                    left.append((i, x[:40]))
    print("remaining 平板 wording:", left or "none")


if __name__ == "__main__":
    main()
