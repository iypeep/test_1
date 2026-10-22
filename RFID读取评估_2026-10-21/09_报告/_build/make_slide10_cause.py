# -*- coding: utf-8 -*-
"""在第 9 页（现网实测基线）之后插入「95% 读 0 的归因」页，并把全篇页码改成 n / 11。

与 make_slide3_merge.py / make_slide5_r2.py / make_slide6_r3.py 同一套路：直接改现成 pptx，
可重复执行（已存在同名页会先删掉再插）。

用法：
    python make_slide10_cause.py                                  # 改 09_报告/RFID读取能力评估_2026-10-22.pptx
    python make_slide10_cause.py --pptx <路径> --out <输出路径>     # 不改原文件
"""
import argparse, os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

_HERE = os.path.dirname(os.path.abspath(__file__))
REP = os.path.dirname(_HERE)
DEFAULT = os.path.join(REP, "RFID读取能力评估_2026-10-22.pptx")

INK, GRAY, LINE = "1A1A1A", "8C8C8C", "D9D9D9"
ACCENT, ACCENT_2, TINT = "3E5DA8", "5C7ED4", "E6EDF9"
RED = "C0504D"
LATIN, EA = "Arial", "微软雅黑"
W, H, MARGIN = 13.333, 7.5, 0.62
TAG = "95% 读 0 的归因"          # 用于幂等判重


# ------------------------------------------------------------------ 基础绘图
def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def textbox(slide, x, y, w, h, text, size=11, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        r = p.add_run(); r.text = line
        set_font(r, size, bold, color)
    return tb


def rich(slide, x, y, w, h, parts, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (t, sz, bold, color) in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        r = p.add_run(); r.text = t
        set_font(r, sz, bold, color)
    return tb


def hline(slide, x, y, w, color=LINE, pt=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = RGBColor.from_string(color)
    ln.line.width = Pt(pt)


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.text_frame.word_wrap = True
    return sh


def _border(cell, color=LINE, pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(pt * 12700))); ln.set("cap", "flat")
        ln.set("cmpd", "sng"); ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "solid")
        tcPr.insert(i, ln)


def table(slide, rows, x, y, w, col_w, row_h=0.34, fs=9, hfs=8.6,
          aligns=None, bold_first_col=False, header_fill="F2F2F2"):
    nrow, ncol = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(row_h * nrow))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    tbl.first_row = False
    tbl.horz_banding = False
    scale = w / sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw * scale)
    for ri in range(nrow):
        tbl.rows[ri].height = Inches(row_h)
        for ci in range(ncol):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.055)
            cell.margin_top = cell.margin_bottom = Inches(0.015)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(header_fill if ri == 0 else "FFFFFF")
            _border(cell)
            p = cell.text_frame.paragraphs[0]
            p.alignment = aligns[ci] if aligns else (PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT)
            p.line_spacing = 1.0
            r = p.add_run(); r.text = str(rows[ri][ci])
            set_font(r, hfs if ri == 0 else fs, (ri == 0) or (bold_first_col and ci == 0))


def stat(slide, x, y, w, num, unit, label, color=ACCENT):
    textbox(slide, x, y, w, 0.5, num, size=26, bold=True, color=color)
    textbox(slide, x, y + 0.46, w, 0.26, unit, size=9.2, color=GRAY)
    textbox(slide, x, y + 0.72, w, 0.55, label, size=9.4, color=INK, spacing=1.15)


# ------------------------------------------------------------------ 幻灯页工具
def slide_title_text(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ""


def delete_slide(prs, idx):
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    rId = items[idx].get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(items[idx])


def move_slide(prs, from_idx, to_idx):
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    el = items[from_idx]
    sldIdLst.remove(el)
    sldIdLst.insert(to_idx, el)


def renumber(prs, total):
    """把每页右下角「n / m」样式页码统一改成 index / total"""
    pat = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
    for i, s in enumerate(prs.slides, start=1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text
            if pat.match(txt):
                p = sh.text_frame.paragraphs[0]
                if p.runs:
                    p.runs[0].text = f"{i} / {total}"
                    for extra in p.runs[1:]:
                        extra.text = ""
                else:
                    p.add_run().text = f"{i} / {total}"


def ensure_footers(prs, total):
    """有的增量脚本会把页脚一起清掉（如 make_slide6_r3 清空非标题形状）——这里补齐"""
    pat = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
    for i, s in enumerate(prs.slides, start=1):
        has_proj = any(sh.has_text_frame and "PALLET" in sh.text_frame.text for sh in s.shapes)
        has_num = any(sh.has_text_frame and pat.match(sh.text_frame.text) for sh in s.shapes)
        if has_proj and has_num:
            continue
        if not has_proj:
            textbox(s, MARGIN, 7.02, 8.0, 0.26, "RFID 读取能力评估 · 7H3AA-PALLET · 2026-10-22",
                    size=8, color=GRAY)
        if not has_num:
            textbox(s, W - MARGIN - 0.8, 7.02, 0.8, 0.26, f"{i} / {total}", size=8, color=GRAY,
                    align=PP_ALIGN.RIGHT)


def replace_text_all(prs, pairs):
    """跨页替换文案（含表格单元格），返回替换计数"""
    n = 0
    def fix(tf):
        nonlocal n
        for p in tf.paragraphs:
            for r in p.runs:
                for a, b in pairs:
                    if a in r.text:
                        r.text = r.text.replace(a, b); n += 1
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                fix(sh.text_frame)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        fix(c.text_frame)
    return n


# ------------------------------------------------------------------ 主流程
def build(pptx_path, out_path):
    prs = Presentation(pptx_path)

    # 幂等：已有同名页则删除
    for i, s in enumerate(prs.slides):
        if TAG in slide_title_text(s):
            delete_slide(prs, i)
            break

    total = len(prs.slides) + 1
    BLANK = prs.slide_layouts[6]
    s = prs.slides.add_slide(BLANK)

    # ---- 页头
    textbox(s, MARGIN, 0.40, W - 2 * MARGIN, 0.55, "95% 读 0 的归因：不是读不到，是绝大多数箱没被盘存",
            size=22, bold=True)
    textbox(s, MARGIN, 0.96, W - 2 * MARGIN, 0.30,
            "同一批 31 天数据拆开看：95% 的箱 0 次读取；命中的 5% 箱平均 18 次，却只读到 1–2 个标签", size=11, color=GRAY)
    hline(s, MARGIN, 1.32, W - 2 * MARGIN)

    # ---- 三个数字
    stat(s, MARGIN, 1.44, 3.72, "0.08 / 6", "31 天 113,667 箱的平均读到标签数",
         "不是“偶尔漏一个”，是绝大多数箱几乎不读")
    stat(s, 4.50, 1.44, 3.72, "18.1 次", "命中箱（5,605 箱）的平均读取事件",
         "一旦被盘存就反复读 —— 双峰，不是均匀漏读", color=ACCENT_2)
    stat(s, 8.40, 1.44, 4.28, "4.9% ≠ 8.6%", "窗口内 ≥1 标签 5,605 箱 vs readStatus=True 9,782 箱",
         "readStatus 判的不是“这个窗口读到任意标签”", color=RED)

    # ---- 左：三种解释
    textbox(s, MARGIN, 2.60, 6.05, 0.26, "三种解释（不互斥）与手上的证据", size=10, bold=True)
    blocks = [
        ("A  触发 / 到位判定",
         "95% 的箱在窗口内 0 次读取；命中的箱却平均 18 次\n⇒ 更像“箱没进入盘存”，不是“弱到读不到”"),
        ("B  readStatus 口径",
         "窗口内 ≥1 标签 4.9%、readStatus=True 8.6%，两组既不等也不互相包含\n⇒ 它判的是指定标签 / 读全条件，需程序方确认"),
        ("C  标签侧（贴标位置 / 遮挡 / 极化）",
         "命中的箱也只读到 1–2 个标签、且反复读同一两个\n⇒ 部分标签被挡住或被箱体极化屏蔽，不是全都弱"),
    ]
    y = 2.90
    for head, body in blocks:
        rect(s, MARGIN, y, 6.05, 1.02, fill="FAFBFC", line="E6E6E6")
        textbox(s, MARGIN + 0.16, y + 0.09, 5.70, 0.24, head, size=9.6, bold=True, color=ACCENT)
        textbox(s, MARGIN + 0.16, y + 0.34, 5.70, 0.62, body, size=8.8, color=INK, spacing=1.2)
        y += 1.10

    # ---- 右：判别实验
    textbox(s, 6.95, 2.60, 5.76, 0.26, "40 分钟就能定性的判别实验（不用改硬件）", size=10, bold=True)
    rows = [
        ["#", "做法", "若结果是这样，则结论"],
        ["1", "箱面额外贴 1 片已知好标签，正常过箱", "读到这片标签但 readStatus 仍 False ⇒ B"],
        ["2", "rfidmon 占会话，同一箱手动触发一次盘存", "能读到 6 个 ⇒ A，与射频能力无关"],
        ["3", "不装金属背壳 / 把标签面对读码器再过一次", "变成 6/6 ⇒ C（遮挡或贴标位置）"],
    ]
    table(s, rows, 6.95, 2.90, 5.76, [0.40, 2.62, 2.74], row_h=0.44, fs=8.8, hfs=8.5,
          aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT])

    rect(s, 6.95, 4.84, 5.76, 1.44, fill=TINT)
    rich(s, 7.13, 4.96, 5.40, 1.26, [
        ("这一步和 P5–P7 的模型预测什么关系", 9.6, True, ACCENT),
        ("模型说的是“布好之后能不能读到 6 个”，前提是箱真的被盘存、标签真的在箱里。", 8.8, False, INK),
        ("现网实测（三面天线 <5 cm 仍 0.08/6、读到数与距离无相关）说明这个前提现在不成立。", 8.8, False, INK),
        ("所以改造顺序必须是：先把“每箱被盘存 + 6 个标签都在可读位置”做实 → 再谈增益/厂家。", 8.8, True, ACCENT),
    ], spacing=1.25)

    # ---- 结论条 + 页脚
    rect(s, MARGIN, 6.42, 0.05, 0.34, fill=ACCENT)
    textbox(s, MARGIN + 0.16, 6.44, 11.5, 0.32,
            "先在 40 分钟内定性 A/B/C；定性之前，任何“换读码器 / 换天线”的收益都无法验证",
            size=11, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, MARGIN, 7.02, 8.0, 0.26, "RFID 读取能力评估 · 7H3AA-PALLET · 2026-10-22", size=8, color=GRAY)
    textbox(s, W - MARGIN - 0.8, 7.02, 0.8, 0.26, f"{total} / {total}", size=8, color=GRAY, align=PP_ALIGN.RIGHT)
    hline(s, MARGIN, 6.94, W - 2 * MARGIN, color="ECECEC")

    # ---- 移到第 9 页之后（0-based 索引 9）
    move_slide(prs, len(prs.slides) - 1, 9)

    # ---- 页码与交叉引用
    renumber(prs, total)
    ensure_footers(prs, total)
    replace_text_all(prs, [("P9–P10", f"P9–P{total}"), ("见 P10", f"见 P{total}")])

    # ---- 讲稿
    sl = prs.slides[9]
    sl.notes_slide.notes_text_frame.text = (
        "这页回答全场一定会被问的那句：既然 P5–P7 的模型说近距离能读全 6 个，为什么现网 95% 读 0？"
        "答案在数据形态上：95% 的箱窗口内 0 次读取，命中的那 5% 箱却平均 18 次读取，而且只读到 1–2 个标签 —— "
        "读取是双峰的，说明多数箱根本没被盘存，而不是射频弱到读不到。"
        "第二点，readStatus=True 的 8.6% 与窗口内真读到标签的 4.9% 不是同一批，说明 readStatus 判的不是“读到任意标签”，语义要找程序方确认。"
        "第三点，命中的箱反复读同一两个标签，说明剩下的标签被位置/遮挡/极化挡住了。"
        "所以下一步不是买读码器，是 40 分钟做三组对照：贴一片已知好标签、rfidmon 手动触发一次、去掉金属背壳再读。"
    )

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    print(f"saved {out_path}  slides={len(prs.slides)}  bytes={os.path.getsize(out_path)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", default=DEFAULT)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    build(a.pptx, a.out or a.pptx)
