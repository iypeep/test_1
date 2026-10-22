# -*- coding: utf-8 -*-
"""重做「R2 模组选型」页：让表格与右侧手册/产品图一一对应（① ② ③），并补手册出处。

解决的问题：
  * 原页上方是 3 款产品图、下方却列了 7 款型号 —— 图文对不上；
  * 手册参数没有出处，评审问“哪来的”答不上。

做法：
  * 表格 7 行，最后加「图」列，只对拿到公开产品图/手册图的 3 款标 ① ② ③；
  * 右侧改成竖排 3 张手册/产品图（figs/hw_strip_r2.png），角标与表格一致；
  * 底部补「手册出处」与两条与实测相关的提醒（含 R3 标定结论）。

用法：
    python make_slide5_r2.py --pptx "E:\\Desktop\\RFID读取能力评估-改动版.pptx"
只改第 5 页（索引 4），可重复执行（按 S5_ 前缀先清后加）。
"""
import os, argparse, shutil, datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

_HERE = os.path.dirname(os.path.abspath(__file__))
REP = os.path.dirname(_HERE)
FIG = os.path.join(REP, "figs")

INK, GRAY, LINE = "1A1A1A", "8C8C8C", "D9D9D9"
ACCENT, TINT, RED, GREEN = "3E5DA8", "E6EDF9", "C0504D", "2E8B57"
LATIN, EA = "Arial", "微软雅黑"
W, MARGIN = 13.333, 0.62
SLIDE_NO = 5

ROWS = [
    ["型号（形态）", "口数 · 功率", "读取率", "灵敏度", "过板 6/6", "图"],
    ["Impinj R700（工业读头）", "4 口 · 33 dBm", "1100 tags/s", "−92 dBm", "● 首选", "—"],
    ["Impinj R420（工业读头）", "4 口 · 32.5 dBm", "1100 tags/s", "−84 dBm", "● 满足", "—"],
    ["RodinBell M-703（E710）", "8 口 · 33 dBm", "900 tags/s", "—", "● 覆盖好", "①"],
    ["Zebra FX9600（一体机）", "4/8 口 · 33 dBm PoE+", "—", "−86 dBm", "● 满足", "—"],
    ["Kathrein RRU 4560 / 4xxx", "4 口 · 30/33 dBm", "—", "—", "● 满足", "②"],
    ["ThingMagic M6e（模块）", "4 口 · 模块", "750 tags/s", "—", "○ 需实测", "—"],
    ["Silion SQM5100（QM100）", "模块 · 需外置 PA", "待定", "−79 dBm", "▲ 先整机实测", "③"],
]


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def textbox(slide, x, y, w, h, text, size=10, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        r = p.add_run(); r.text = line
        set_font(r, size, bold, color)
    return tb


def rich(slide, x, y, w, h, parts, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (t, sz, bold, color) in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        r = p.add_run(); r.text = t
        set_font(r, sz, bold, color)
    return tb


def _border(cell, color=LINE, pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(pt * 12700))); ln.set("cap", "flat")
        ln.set("cmpd", "sng"); ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "solid")
        tcPr.insert(i, ln)


def table(slide, rows, x, y, w, col_w, row_h=0.3, fs=8.8, hfs=8.4,
          aligns=None, bold_first_col=False, header_fill="F2F2F2", highlight_rows=(), highlight_fill=TINT):
    nrow, ncol = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(row_h * nrow))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    tbl.first_row = False
    tbl.horz_banding = False
    scale = w / sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw * scale)
    for ri in range(nrow):
        tbl.rows[ri].height = Inches(row_h)
        for ci in range(ncol):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                header_fill if ri == 0 else (highlight_fill if ri in highlight_rows else "FFFFFF"))
            _border(cell)
            p = cell.text_frame.paragraphs[0]
            p.alignment = aligns[ci] if aligns else (PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT)
            p.line_spacing = 1.0
            r = p.add_run(); r.text = str(rows[ri][ci])
            set_font(r, hfs if ri == 0 else fs, (ri == 0) or (bold_first_col and ci == 0))
    return shp


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", default="")
    a = ap.parse_args()
    src, out = a.pptx, (a.out or a.pptx)
    if out == src:
        bak = os.path.splitext(src)[0] + ".bak-r2-" + datetime.datetime.now().strftime("%H%M%S") + ".pptx"
        shutil.copy2(src, bak); print("backup:", bak)

    prs = Presentation(src)
    s = prs.slides[SLIDE_NO - 1]

    # 清旧：上一版脚本加的 S5_*，以及原页的图/表（图文不一致的根源）
    for sh in list(s.shapes):
        drop = sh.name.startswith("S5_") or sh.name in ("Picture 9", "Table 11")
        if drop:
            sh._element.getparent().remove(sh._element)

    # 标题副行
    tb = textbox(s, MARGIN, 0.78, W - 2 * MARGIN, 0.26,
                 "判据＝一次过板读全 6 台；下表 7 款候选，只有 3 款拿到公开产品图 / 手册图（①②③），其余参数只能引手册",
                 size=9.5, color=GRAY)
    tb.name = "S5_sub"

    # 表格
    t = table(s, ROWS, MARGIN, 1.18, 8.35,
              [2.6, 1.95, 1.3, 1.15, 1.0, 0.45], row_h=0.40, fs=8.6, hfs=8.2,
              bold_first_col=True,
              aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER],
              highlight_rows=(3, 5, 7))
    t.name = "S5_table"

    # 右侧手册/产品图（竖排，角标 ①②③）
    p = os.path.join(FIG, "hw_strip_r2.png")
    iw, ih = Image.open(p).size
    maxw, maxh = 3.45, 5.05
    sc = min(maxw / iw, maxh / ih)
    pic = s.shapes.add_picture(p, Inches(9.22), Inches(1.18), Inches(iw * sc), Inches(ih * sc))
    pic.name = "S5_strip"

    # 说明块
    rich(s, MARGIN, 4.60, 8.35, 1.85, [
        ("选型不能只看功率", 9.8, True, ACCENT),
        ("读全判据：6 个标签 ×≥2 次盘存 ÷ 0.5 s ≥ 24 tags/s；手册 tags/s 是理想值，箱内按 7 折估。", 8.8, False, INK),
        ("国产 QM100 类片内 PA 仅 4 dBm，整机 30 dBm 靠外置功放 —— 只能按整机实测比，不能拿芯片指标比。", 8.8, False, INK),
        ("▲ 与 R3 标定一致：现网三面天线已贴到 <5 cm，实测仍只有 0.08/6 ⇒ 换模块的收益必须先实测验证，"
         "否则换了也读不全（瓶颈更可能在标签朝向/金属屏蔽/读码窗口）。", 8.8, True, RED),
        ("手册出处：Kathrein RRU 4000 系列 datasheet（背板 ANT1–4）；RodinBell M-703 / E710 产品页；"
         "Silion SQM5100（RFID-CM100）装置图；Impinj Fixed Reader Portfolio、Zebra FX9600 Spec Sheet（无公开产品图）。",
         7.6, False, GRAY),
    ], spacing=1.22)

    # 底部结论
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(6.60), Inches(0.05), Inches(0.30))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(ACCENT); bar.line.fill.background()
    bar.shadow.inherit = False
    bar.name = "S5_bar"
    tb = textbox(s, MARGIN + 0.16, 6.61, W - 2 * MARGIN - 0.2, 0.28,
                 "选型顺序：先按「天线口数/覆盖 + 灵敏度」筛，再按「整机实测 tags/s」定；R700 / R420 / M-703 为第一梯队，国产模块先做整机验证。",
                 size=9.6, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb.name = "S5_take"

    prs.save(out)
    print("saved", out, os.path.getsize(out), "bytes; slides =", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
