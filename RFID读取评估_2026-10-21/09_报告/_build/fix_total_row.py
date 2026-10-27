# -*- coding: utf-8 -*-
"""合计行改为「整列求和」口径（用户选方案 2）：177,306 / 6,869 / 0.30%

说明：用户原话给的是 5,947 / 2.94%，但按表内七行实际求和是
    箱数 177,306 ｜ 读到 ≥1 箱 6,869（3.87%）
    读全 6 个按各行百分比加权 = 535 箱 = 0.30%
所以取"列能对上"的值；合计行左侧标注保持「四个在产组合」会自相矛盾，改成「全部在产组合（含推算）」。

用法：py -3.12 fix_total_row.py --pptx <deck>.pptx --out <新.pptx>
"""
import os, argparse
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

LATIN, EA, INK = "Arial", "微软雅黑", "1A1A1A"
TOTAL = {1: "全部在产组合（含推算）", 4: "177,306", 5: "6,869", 6: "0.30%"}


def set_font(run, size=9.4, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.pptx)
    t = None
    for sl in prs.slides:
        if any(sh.name == "TOC_title" for sh in sl.shapes):
            continue
        for sh in sl.shapes:
            if getattr(sh, "has_table", False) and sh.has_table and sh.table.cell(0, 0).text.strip() == "HQ代码":
                t = sh.table; break
        if t is not None:
            break
    if t is None:
        raise SystemExit("找不到机型统计表")
    last = len(t.rows) - 1
    for ci, val in TOTAL.items():
        cell = t.cell(last, ci)
        cell.text_frame.paragraphs[0].clear()
        r = cell.text_frame.paragraphs[0].add_run()
        r.text = val
        set_font(r, 9.4, ci in (4, 5, 6))
    prs.save(a.out)
    print("saved", a.out)
    for ri in range(len(t.rows)):
        print("  ", [t.cell(ri, c).text for c in range(len(t.columns))])


if __name__ == "__main__":
    main()
