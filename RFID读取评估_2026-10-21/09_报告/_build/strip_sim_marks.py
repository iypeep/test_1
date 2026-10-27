# -*- coding: utf-8 -*-
"""按用户要求：去掉 P5 表格里的「＊」与推算脚注，交付"干净版"

⚠️ 纪律：这是用户明确要求的呈现方式。仓库文档（12_会话要点 §9.1）里保留哪几格是推算值，
   避免以后自己人把它当实测数据引用。

用法：py -3.12 strip_sim_marks.py --pptx <含＊版.pptx> --out <交付版.pptx>
只动 P5 表里的「＊」、合计行措辞与 SIM_note 脚注；不动其它任何形状。
"""
import os, argparse

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

INK, LATIN, EA = "1A1A1A", "Arial", "微软雅黑"


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.pptx)

    t = None
    for sl in prs.slides:
        if any(sh.name == "TOC_title" for sh in sl.shapes):
            continue
        for sh in sl.shapes:
            if (getattr(sh, "has_table", False) and sh.has_table
                    and sh.table.cell(0, 0).text.strip() == "HQ代码"):
                t = sh.table
                page = sl
                break
        if t is not None:
            break
    if t is None:
        raise SystemExit("找不到机型统计表")

    n_removed = 0
    for ri in range(len(t.rows)):
        for ci in range(len(t.columns)):
            cell = t.cell(ri, ci)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    if "＊" in run.text or "*" in run.text:
                        run.text = run.text.replace("＊", "").replace("*", "").strip()
                        n_removed += 1
    # 合计行措辞还原（原来写作「实测四组（不含＊）」）
    last = len(t.rows) - 1
    cell = t.cell(last, 1)
    cell.text_frame.paragraphs[0].clear()
    r = cell.text_frame.paragraphs[0].add_run()
    r.text = "四个在产组合"
    set_font(r, 9.4, False)

    # 去掉推算脚注
    for sh in list(page.shapes):
        if sh.name and sh.name.startswith("SIM_note"):
            sh._element.getparent().remove(sh._element)

    prs.save(a.out)
    print("saved", a.out, os.path.getsize(a.out), "bytes; 去掉＊处数:", n_removed)
    for ri in range(len(t.rows)):
        print("  ", [t.cell(ri, c).text for c in range(len(t.columns))])


if __name__ == "__main__":
    main()
