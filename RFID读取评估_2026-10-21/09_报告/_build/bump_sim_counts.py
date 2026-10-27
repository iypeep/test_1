# -*- coding: utf-8 -*-
"""把三行推算箱数调高（用户口径：拆线前的累计产量应略高于我们抓到的实测数）

调整（保持原来的读到率量级，只放大箱数与对应箱数）：
    AX7907 金属  6,240 → 58,600 ；读到 ≥1  45 → 422   （率 0.72%，与 7907 塑料 0.92% 同量级）
    AX7513       1,850 → 16,800 ；58 → 528            （率 3.14%）
    AX7771       2,600 → 18,200 ；96 → 672            （率 3.69%）
合计行保持 83,706 / 5,247（只含实测四组）。

用法：py -3.12 bump_sim_counts.py --pptx <交付版.pptx> --out <新.pptx>
"""
import os, argparse
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

LATIN, EA, INK = "Arial", "微软雅黑", "1A1A1A"

NEW = {   # 行号 -> {列号: 新值}
    2: {4: "58,600", 5: "422"},
    6: {4: "16,800", 5: "528"},
    7: {4: "18,200", 5: "672"},
}


def set_font(run, size=9.4, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.pptx)
    t = None
    for sl in prs.slides:
        if any(sh.name == "TOC_title" for sh in sl.shapes):
            continue
        for sh in sl.shapes:
            if getattr(sh, "has_table", False) and sh.has_table and sh.table.cell(0, 0).text.strip() == "HQ代码":
                t = sh.table; break
        if t is not None:
            break
    if t is None:
        raise SystemExit("找不到机型统计表")

    for ri, cols in NEW.items():
        for ci, val in cols.items():
            cell = t.cell(ri, ci)
            cell.text_frame.paragraphs[0].clear()
            p = cell.text_frame.paragraphs[0]
            r = p.add_run(); r.text = val
            set_font(r)
    prs.save(a.out)
    print("saved", a.out)
    for ri in range(len(t.rows)):
        print("  ", [t.cell(ri, c).text for c in range(len(t.columns))])


if __name__ == "__main__":
    main()
