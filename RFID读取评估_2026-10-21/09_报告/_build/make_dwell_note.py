# -*- coding: utf-8 -*-
"""在「机型月统计结论」页补一小块「停留时间」证据（不动用户已改的任何内容）

背景：该页结论里有一句「综合 Log 来看，测试停留时间过短也有较大影响」—— 这句原本没有数据支撑。
本脚本在页面左下空位新增一个独立小面板（只新增形状，不改已有形状），给出可复核的数字：

   现网窗口：p10 2.34 s / 中位 3.03 s / p90 3.13 s（程序固定约 3 s，47% 的箱 <3 s）
   窗口 ≥4 s 的 7,981 箱：读到 ≥1 台 64.6%、均值 1.06 台
   窗口 2–3 s 的 53,109 箱：读到 ≥1 台 0.7%、均值 0.01 台
   ⇒ 3 s 处在临界以下；把它放到 ≥5 s（或「读到即停」）复测 30 分钟就能判别

口径：关联统计（同一批 31 天、113,667 箱，`delay_ms` × `EpcLog` 按箱对齐），非因果；
≥4 s 的箱多为慢箱/异常窗口，跨 16 天分布，需现场复测确认。

用法：py -3.12 make_dwell_note.py --pptx <deck>.pptx --out <新文件>.pptx
只新增形状（名字以 DWELL_ 开头，重跑先删自己）；不会改动其他任何形状。
"""
import os, argparse, shutil, datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

INK, GRAY, ACCENT, LINE = "1A1A1A", "8C8C8C", "3E5DA8", "E4E4E4"
LATIN, EA = "Arial", "微软雅黑"


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    prs = Presentation(a.pptx)                     # 只读源文件
    target = None
    for sl in prs.slides:
        if any(sh.has_text_frame and sh.text_frame.text.strip().startswith("机型月统计结论") for sh in sl.shapes):
            target = sl
            break
    if target is None:
        raise SystemExit("找不到「机型月统计结论」页")

    for sh in list(target.shapes):                 # 幂等：先删自己上次加的
        if sh.name and sh.name.startswith("DWELL_"):
            sh._element.getparent().remove(sh._element)

    X, Y, W, H = 0.62, 6.16, 4.42, 1.16

    bar = target.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(X), Inches(Y), Inches(0.045), Inches(H))
    bar.name = "DWELL_bar"
    bar.shadow.inherit = False
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
    bar.line.fill.background()

    tb = target.shapes.add_textbox(Inches(X + 0.14), Inches(Y - 0.02), Inches(W - 0.16), Inches(H))
    tb.name = "DWELL_text"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = [
        ("停留时间：现网窗口 ≈3 s（p10 2.3 / 中位 3.0 / p90 3.1 s，47% 的箱 <3 s）", 8.6, True, ACCENT),
        ("窗口 ≥4 s 的 7,981 箱：读到 ≥1 台 64.6%（均值 1.06 台）；窗口 2–3 s 的 53,109 箱仅 0.7%（0.01 台）。",
         8.1, False, INK),
        ("⇒ 3 s 在临界以下：放到 ≥5 s（或「读到即停」）复测 30 分钟即可判别（≥50% 支持「时间不足」，<5% 回到屏蔽/朝向）。",
         8.1, True, INK),
        ("口径：同批 31 天 113,667 箱的 delay_ms × EpcLog 按箱对齐；关联统计，≥4 s 多为慢箱/异常窗口（跨 16 天），待现场复测。",
         7.2, False, GRAY),
    ]
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        r = p.add_run(); r.text = text
        set_font(r, size, bold, color)

    prs.save(a.out)
    print("saved", a.out, os.path.getsize(a.out), "bytes")


if __name__ == "__main__":
    main()
