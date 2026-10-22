# -*- coding: utf-8 -*-
"""调整 P3（读取方向测试）：把「中箱读取测试」表的表头 上/下/左/右/前/后
与 P2「六种放置方式」的 A1–A6 / 面代号一一对上。

做法（可重复执行）：
  1) 表 S3_boxtable 里插一行「P2 姿态」（A2/A1/A5/A6/A3/A4 + 面代号），旧的先删；
  2) 表上方加一行图例，说明表头 = 该面朝下、读码器固定在下且朝向不变；
  3) 微调行高与 y 坐标，保证不越界；
  4) 现象里点明「上/下（A2/A1）最好」。

用法：
    python make_p3_boxface.py --pptx <base.pptx> [--out <out.pptx>] [--slide-no 3]
"""
import argparse, copy, datetime, os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

INK, GRAY, ACCENT, TINT = "1A1A1A", "8C8C8C", "3E5DA8", "E6EDF9"
LATIN, EA = "Arial", "微软雅黑"

# 表头列 → P2 的 3D 姿态（该面朝下）
FACE_MAP = [("A2", "T-TOP"), ("A1", "B-BOTTOM"), ("A5", "S-LEFT"),
            ("A6", "S-RIGHT"), ("A3", "F-FRONT"), ("A4", "F-BACK")]
ROW_LABEL = "P2 姿态"
LEGEND = ("表头 ＝ P2「六种放置方式」里该面朝下的姿态：上 A2·T-TOP　下 A1·B-BOTTOM　左 A5·S-LEFT　"
          "右 A6·S-RIGHT　前 A3·F-FRONT　后 A4·F-BACK（读码器固定在下、朝向不变）")


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def set_cell(cell, lines, size=8.2, bold=True, color=INK, fill=None, align=PP_ALIGN.CENTER):
    if fill:
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.03)
    cell.margin_top = cell.margin_bottom = Inches(0.01)
    tf = cell.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, t in enumerate(lines):
        p = first if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 0.92
        r = p.add_run(); r.text = t
        set_font(r, size, bold, color)


def find_slide(prs, keyword, prefer=None):
    if prefer:
        s = prs.slides[prefer - 1]
        for sh in s.shapes:
            if sh.has_text_frame and keyword in sh.text_frame.text:
                return s
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and keyword in sh.text_frame.text and sh.top < Inches(1.0):
                return s
    raise SystemExit(f"没找到标题含「{keyword}」的页")


def get_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", default="")
    ap.add_argument("--slide-no", type=int, default=3)
    a = ap.parse_args()
    src, out = a.pptx, (a.out or a.pptx)
    if out == src:
        bak = os.path.splitext(src)[0] + ".bak-p3-" + datetime.datetime.now().strftime("%H%M%S") + ".pptx"
        shutil.copy2(src, bak); print("backup:", bak)

    prs = Presentation(src)
    s = find_slide(prs, "读取方向测试", a.slide_no)
    tbl_sh = get_shape(s, "S3_boxtable")
    if tbl_sh is None:
        raise SystemExit("P3 上找不到 S3_boxtable")
    tbl = tbl_sh.table

    # ---- 1) 幂等：删掉已有的「P2 姿态」行
    for ri, row in enumerate(list(tbl.rows)):
        if row.cells[0].text.strip() == ROW_LABEL:
            tr = row._tr
            tr.getparent().remove(tr)
            print("已移除旧的 P2 姿态行")

    # ---- 2) 复制表头行做样式模板，插到第 2 行
    trs = tbl._tbl.findall(qn("a:tr"))
    new_tr = copy.deepcopy(trs[0])
    trs[0].addnext(new_tr)
    tbl2 = tbl_sh.table                     # 重新取，保证行对象已刷新
    new_row = tbl2.rows[1]
    set_cell(new_row.cells[0], [ROW_LABEL], size=8.0, bold=True, color=ACCENT, fill=TINT,
             align=PP_ALIGN.LEFT)
    for ci, (code, face) in enumerate(FACE_MAP, start=1):
        set_cell(new_row.cells[ci], [code, face], size=7.8, bold=True, color=ACCENT, fill=TINT)

    # ---- 3) 行高：表头0.26 / P2行0.32 / 其余 0.25
    tbl2.rows[0].height = Inches(0.26)
    tbl2.rows[1].height = Inches(0.32)
    for ri in range(2, len(tbl2.rows)):
        tbl2.rows[ri].height = Inches(0.25)

    # ---- 4) 标题上移一线 + 图例 + 表位置
    title = get_shape(s, "S3_boxtitle")
    if title is not None:
        title.top = Inches(1.36)
    legend = get_shape(s, "S3_legend")
    if legend is None:
        legend = s.shapes.add_textbox(Inches(7.15), Inches(1.62), Inches(5.56), Inches(0.34))
        legend.name = "S3_legend"
    tf = legend.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    p0.line_spacing = 1.15
    r = p0.add_run(); r.text = LEGEND
    set_font(r, 7.4, False, GRAY)
    legend.top, legend.height = Inches(1.62), Inches(0.34)

    tbl_sh.top = Inches(1.94)

    # ---- 5) 现象文字：点明上/下＝A2/A1
    boxtext = get_shape(s, "S3_boxtext")
    if boxtext is not None:
        boxtext.top = Inches(3.78)
        for p in boxtext.text_frame.paragraphs:
            for r in p.runs:
                if "姿态差异明显" in r.text:
                    r.text = r.text.replace("姿态差异明显：", "姿态差异明显（上/下＝A2/A1 最好）：")
    src_note = get_shape(s, "S3_source")
    if src_note is not None:
        src_note.top = Inches(5.92)

    # ---- 6) 清掉错位的旧页码（如 4 / 10：页数已变，且只有 3 页有）
    import re as _re
    pat = _re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
    n = 0
    for sl in prs.slides:
        for sh in list(sl.shapes):
            if sh.has_text_frame and pat.match(sh.text_frame.text):
                sh._element.getparent().remove(sh._element); n += 1
    if n:
        print(f"移除错位页码 {n} 处")

    prs.save(out)
    print(f"saved {out}  slides={len(prs.slides)}  bytes={os.path.getsize(out)}")


if __name__ == "__main__":
    main()
