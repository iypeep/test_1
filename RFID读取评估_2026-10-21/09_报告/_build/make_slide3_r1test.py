# -*- coding: utf-8 -*-
"""把 2026-10-19 中箱测试结果表放进 P3（R1 放置方式页），
表头 上/下/左/右/前/后 用 R1 的六面定义（A1–A6 / 面代号）说明。

做法：与 make_slide6_r3.py 一致 —— 只重建目标页的正文（保留标题与分割线），可重复执行。

用法：
    python make_slide3_r1test.py                                  # 改 09_报告/RFID读取能力评估_2026-10-22.pptx 的第 3 页
    python make_slide3_r1test.py --slide-no 3 --pptx <path> [--out <path>]
"""
import argparse, os, re, datetime, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

_HERE = os.path.dirname(os.path.abspath(__file__))
REP = os.path.dirname(_HERE)
ROOT = os.path.dirname(REP)
FIG = os.path.join(REP, "figs")
DEFAULT = os.path.join(REP, "RFID读取能力评估_2026-10-22.pptx")

INK, GRAY, LINE = "1A1A1A", "8C8C8C", "D9D9D9"
ACCENT, ACCENT_2, TINT, RED = "3E5DA8", "5C7ED4", "E6EDF9", "C0504D"
LATIN, EA = "Arial", "微软雅黑"
W, H, MARGIN = 13.333, 7.5, 0.62

# 2026-10-19 中箱读取测试原始数据（读出的平板个数 / 6；None = 空白 = 未读到）
COLS = [("上", "T-TOP", "A2"), ("下", "B-BOTTOM", "A1"), ("左", "S-LEFT", "A5"),
        ("右", "S-RIGHT", "A6"), ("前", "F-FRONT", "A3"), ("后", "F-BACK", "A4")]
ROWS = [
    ("大平板 1", [1, None, None, 1, None, None]),
    ("大平板 2", [1, 1, 1, 1, None, None]),
    ("大平板 3", [1, 1, None, None, None, None]),
    ("大平板 4", [None, 1, None, None, None, None]),
    ("小平板 1", [1, 1, None, None, None, None]),
    ("小平板 2", [2, None, 2, None, None, None]),
    ("小平板 3", [None, None, None, None, 1, None]),
    ("小平板 4", [None, 1, None, None, None, None]),
]


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def textbox(slide, x, y, w, h, text, size=11, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        r = p.add_run(); r.text = line
        set_font(r, size, bold, color)
    return tb


def rich(slide, x, y, w, h, parts, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (t, sz, bold, color) in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        r = p.add_run(); r.text = t
        set_font(r, sz, bold, color)
    return tb


def hline(slide, x, y, w, color=LINE, pt=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = RGBColor.from_string(color)
    ln.line.width = Pt(pt)


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.text_frame.word_wrap = True
    return sh


def pic(slide, path, x, y, maxw, maxh):
    iw, ih = Image.open(path).size
    s = min(maxw / iw, maxh / ih)
    slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(iw * s), height=Inches(ih * s))
    return iw * s, ih * s


def _border(cell, color=LINE, pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(pt * 12700))); ln.set("cap", "flat")
        ln.set("cmpd", "sng"); ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "solid")
        tcPr.insert(i, ln)


def cell_text(cell, lines, size, bold=False, color=INK, align=PP_ALIGN.CENTER):
    """支持多行（用 \n 分隔的 list/tuple）"""
    if isinstance(lines, str):
        lines = lines.split("\n")
    tf = cell.text_frame
    tf.word_wrap = True
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 0.95
        r = p.add_run(); r.text = t
        set_font(r, size, bold, color)


def build_table(slide, x, y, w, row_h, header_h):
    ncol = 1 + len(COLS)
    nrow = 1 + len(ROWS) + 1                     # 表头 + 8 轮 + 合计
    shp = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w),
                                 Inches(header_h + row_h * (nrow - 1)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    tbl.first_row = False
    tbl.horz_banding = False

    widths = [1.42] + [(w - 1.42) / len(COLS)] * len(COLS)
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for ri in range(1, nrow):
        tbl.rows[ri].height = Inches(row_h)

    # 表头：上 / T-TOP / A2（三行）
    hdr = ["箱体哪个面朝下\n（接触底部读码器）"] + [f"{c}\n{face}\n{code}" for c, face, code in COLS]
    for ci, txt in enumerate(hdr):
        cell = tbl.cell(0, ci)
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string("F2F2F2")
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.margin_top = cell.margin_bottom = Inches(0.01)
        _border(cell)
        cell_text(cell, txt, 8.2, bold=True,
                  align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)

    # 数据行
    for ri, (label, vals) in enumerate(ROWS, start=1):
        cell = tbl.cell(ri, 0)
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.03)
        cell.margin_top = cell.margin_bottom = Inches(0.01)
        _border(cell)
        cell_text(cell, label, 8.8, bold=True, align=PP_ALIGN.LEFT)
        for ci, v in enumerate(vals, start=1):
            c = tbl.cell(ri, ci)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string("FFFFFF" if v is None else TINT)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_top = c.margin_bottom = Inches(0.01)
            _border(c)
            cell_text(c, "" if v is None else str(v), 9.4, bold=(v is not None),
                      color=INK if v is not None else GRAY)

    # 合计行
    totals = [sum(1 for _, v in ROWS if v[i] is not None) for i in range(len(COLS))]
    ri = nrow - 1
    cell = tbl.cell(ri, 0)
    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string("F2F2F2")
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.06)
    _border(cell)
    cell_text(cell, "读到次数（处）", 8.6, bold=True, align=PP_ALIGN.LEFT)
    for ci, v in enumerate(totals, start=1):
        c = tbl.cell(ri, ci)
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string("F2F2F2")
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        _border(c)
        cell_text(c, str(v), 9.2, bold=True, color=RED if v == 0 else INK)
    return shp


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", default=DEFAULT)
    ap.add_argument("--out", default="")
    ap.add_argument("--slide-no", type=int, default=3, help="目标页（1-based），本仓库底稿 R1 是第 3 页")
    a = ap.parse_args()
    src, out = a.pptx, (a.out or a.pptx)
    if out == src:
        bak = os.path.splitext(src)[0] + ".bak-r1-" + datetime.datetime.now().strftime("%H%M%S") + ".pptx"
        shutil.copy2(src, bak); print("backup:", bak)

    prs = Presentation(src)
    s = prs.slides[a.slide_no - 1]

    # 保留标题与分割线，其余全部清掉（可重复执行）
    for sh in list(s.shapes):
        keep_title = sh.has_text_frame and "R1" in sh.text_frame.text and sh.top is not None and sh.top < Inches(0.9)
        is_line = str(sh.shape_type).startswith("LINE") or sh.name.startswith("Connector")
        if keep_title or is_line:
            continue
        sh._element.getparent().remove(sh._element)

    # 副标题
    textbox(s, MARGIN, 0.96, W - 2 * MARGIN, 0.30,
            "六面分别作底面（A1–A6）；右下表 = 2026-10-19 中箱测试原始读数（每格 = 该姿态读出的平板数 / 6，空白 = 未读到）",
            size=10.5, color=GRAY)

    # 左：六宫格 3D
    pic(s, os.path.join(FIG, "fig4_3d_placements.png"), MARGIN, 1.36, 6.30, 4.10)

    # 右上：一行图例（表头 ↔ 面定义 ↔ A 编号）
    textbox(s, 7.10, 1.36, 5.62, 0.24, "表头 ＝ 箱体哪个面朝下接触底部读码器（对应 A1–A6）",
            size=9.4, bold=True, color=ACCENT)
    textbox(s, 7.10, 1.62, 5.62, 0.42,
            "上 = T-TOP (A2，倒置)　下 = B-BOTTOM (A1，正立)　左 = S-LEFT (A5，侧躺)\n"
            "右 = S-RIGHT (A6，侧躺)　前 = F-FRONT (A3，侧放)　后 = F-BACK (A4，侧放)",
            size=8.4, color=GRAY, spacing=1.25)

    # 右下：测试结果表
    tbl2 = build_table(s, 7.10, 2.10, 5.62, 0.24, 0.46)

    # 表下：现象三条
    rich(s, 7.10, 4.92, 5.62, 1.10, [
        ("三条现象", 9.4, True, ACCENT),
        ("① 单次最多 2/6（目标 6），多数 1 个，不少姿态 0 个。", 8.6, False, INK),
        ("② 上/下各 5 处读到，左/右各 2 处；前 1 处、后 0 处。", 8.6, False, INK),
        ("③ 大平板单次最多 1 个，小平板最多 2 个 → 大件金属多、遮挡更重。", 8.6, False, INK),
    ], spacing=1.25)

    # 左下：数据来源
    textbox(s, MARGIN, 5.50, 6.30, 0.62,
            "数据来源：`RFID中箱读取测试报告_20261019.pdf`（2026-10-19）—— 读取工位在底面/侧面/顶面均布感应器，"
            "箱子以不同面朝下接触底部读码器，共 8 轮（大平板 4 + 小平板 4）。",
            size=8.2, color=GRAY, spacing=1.2)

    # 结论条
    rect(s, MARGIN, 6.36, 0.05, 0.34, fill=ACCENT)
    textbox(s, MARGIN + 0.16, 6.38, 11.4, 0.32,
            "六面实测：前/后朝下基本读不到、上/下最好也只有 1–2/6 —— 与 P5「距离无相关」一致，卡点是金属平板屏蔽与箱内互相遮挡",
            size=10.6, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    prs.save(out)
    print(f"saved {out}  slides={len(prs.slides)}  bytes={os.path.getsize(out)}")


if __name__ == "__main__":
    main()
