# -*- coding: utf-8 -*-
"""按各页现有标题重排目录（只改目录页，不动任何页面的标题）

规则：
  · 取每页「顶部 y<1.2"」的第一段文字作为标题（封面 / Back Up 分隔页没有，自动跳过）；
  · 目录页本身跳过；
  · 行距 0.50"、12pt；标题超过 26 字自动降到 11pt；
  · 目录条目文案＝页面标题原文，不改编。

用法：py -3.12 sync_toc.py --pptx <deck>.pptx --out <新.pptx>
"""
import os, argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

INK, GRAY, ACCENT, LATIN, EA = "1A1A1A", "8C8C8C", "3E5DA8", "Arial", "微软雅黑"
W, MARGIN = 13.333, 0.62


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, LATIN
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", EA if tag == "a:ea" else LATIN)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    set_font(r, size, bold, color)
    if name:
        tb.name = name
    return tb


def hline(slide, x, y, w, color="EFEFEF", pt=0.75, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(pt)
    if name:
        c.name = name
    return c


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.pptx)

    toc = None
    for sl in prs.slides:
        if any(sh.name == "TOC_title" for sh in sl.shapes):
            toc = sl
            break
    if toc is None:
        raise SystemExit("找不到目录页（TOC_title）")

    items = []
    for idx, sl in enumerate(prs.slides, start=1):
        if sl is toc:
            continue
        title = ""
        for sh in sl.shapes:
            if (sh.has_text_frame and sh.top is not None
                    and sh.top < Inches(1.2) and sh.text_frame.text.strip()):
                title = sh.text_frame.text.strip().split("\n")[0].strip()
                break
        if title:
            items.append((title, idx))
    print("目录条目：")
    for t, pg in items:
        print(f"   {pg:>2}  {t}")

    for sh in list(toc.shapes):
        nm = sh.name or ""
        if nm.startswith("TOC_") or str(sh.shape_type).startswith("LINE") or nm.startswith("Straight Connector"):
            sh._element.getparent().remove(sh._element)

    textbox(toc, MARGIN, 0.55, 6.0, 0.55, "目录", 26, True, INK, name="TOC_title")
    hline(toc, MARGIN, 1.30, W - 2 * MARGIN, color="D9D9D9", name="TOC_rule")
    y, step = 1.50, 0.50
    for i, (title, page) in enumerate(items, start=1):
        size = 12 if len(title) <= 26 else 11
        textbox(toc, MARGIN + 0.30, y, 0.55, 0.28, f"{i}", 10.5, True, ACCENT, name=f"TOC_n{i}")
        textbox(toc, MARGIN + 0.85, y - 0.03, 9.30, 0.32, title, size, False, INK, name=f"TOC_t{i}")
        textbox(toc, MARGIN + 10.30, y - 0.03, 1.6, 0.32, str(page), 11, False, GRAY, name=f"TOC_p{i}")
        hline(toc, MARGIN + 0.30, y + 0.44, W - 2 * MARGIN - 0.60, name=f"TOC_l{i}")
        y += step

    prs.save(a.out)
    print("saved", a.out, os.path.getsize(a.out))


if __name__ == "__main__":
    main()
